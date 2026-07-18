"""AC-P5-2: xlsx_builder produces workbook with:
  * 6 techtrade base sheets (Recommendations, Levels, Reasoning, Orders,
    Fills, Summary — verified from techtrade excel_export.SHEET_ORDER)
  * 3 intraday sheets (IntradayFills, Vetoes, PerSymbolPnL) — ours

Extra-missing path: build_workbook raises XLSXUnavailable when
openbb-techtrade is uninstalled, and report(format='xlsx' | 'all')
demotes it to a warning rather than crashing.

Note on the base-sheet name: techtrade already has a "Fills" sheet, so
our intraday one is "IntradayFills" to avoid collision.
"""

from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal

import pytest


@pytest.fixture(autouse=True)
def _allow_absolute_export_paths(monkeypatch):
    """xlsx_builder → techtrade excel_export rejects absolute paths by
    default (security guard). Pytest's tmp_path is always absolute, so
    every test in this module needs the documented escape hatch."""
    monkeypatch.setenv("TECHTRADE_EXPORT_ALLOW_ABSOLUTE", "1")


# Techtrade base workbook sheets — verified from
# openbb_platform/extensions/techtrade/openbb_techtrade/reporting/excel_export.py
# (SHEET_ORDER at line 74).
_EXPECTED_TECHTRADE_SHEETS = {
    "Recommendations", "Levels", "Reasoning", "Orders", "Fills", "Summary",
}
_EXPECTED_INTRADAY_SHEETS = {"IntradayFills", "Vetoes", "PerSymbolPnL"}


class TestXLSXBuilderIntegration:
    """AC-P5-2 real assertion: techtrade base sheets AND our intraday sheets
    are all present. Fails green against a stub that skips techtrade."""

    def test_produces_workbook_with_all_expected_sheets(self, tmp_path):
        pytest.importorskip("openpyxl")
        pytest.importorskip("openbb_techtrade")

        from openbb_fmp_trading.models.report import SessionMetrics
        from openbb_fmp_trading.reporting.xlsx_builder import build_workbook

        output = tmp_path / "end_of_day.xlsx"
        metrics = SessionMetrics(realized_pnl=Decimal("100"))
        build_workbook("s20260713", events=[], metrics=metrics, output_path=output)

        assert output.exists()
        from openpyxl import load_workbook
        wb = load_workbook(output)
        sheet_names = set(wb.sheetnames)

        # Review finding #1: the AC requires the 6 techtrade sheets.
        # This assertion fails if _call_techtrade_export is a stub.
        missing_base = _EXPECTED_TECHTRADE_SHEETS - sheet_names
        assert not missing_base, (
            f"Missing techtrade base sheets: {missing_base}. If "
            f"_call_techtrade_export is stubbed, AC-P5-2 (6 techtrade + "
            f"3 intraday sheets, per design-spec §4.3) is not satisfied. "
            f"Do not close AC-P5-2 on this state."
        )

        # Our 3 intraday sheets MUST be present
        missing_intraday = _EXPECTED_INTRADAY_SHEETS - sheet_names
        assert not missing_intraday, (
            f"Missing intraday sheets: {missing_intraday}"
        )

    def test_intraday_sheets_have_column_headers(self, tmp_path):
        """Sanity: appended sheets have the header row we documented."""
        pytest.importorskip("openpyxl")
        pytest.importorskip("openbb_techtrade")

        from openbb_fmp_trading.models.report import SessionMetrics
        from openbb_fmp_trading.reporting.xlsx_builder import build_workbook

        output = tmp_path / "end_of_day.xlsx"
        build_workbook(
            "s", events=[],
            metrics=SessionMetrics(realized_pnl=Decimal("0")),
            output_path=output,
        )

        from openpyxl import load_workbook
        wb = load_workbook(output)
        assert wb["IntradayFills"][1][0].value == "ts"
        assert wb["Vetoes"][1][0].value == "ts"
        assert wb["PerSymbolPnL"][1][0].value == "symbol"


class TestXLSXExtraHandling:
    def test_extra_missing_raises_xlsxunavailable(self, monkeypatch):
        """When techtrade isn't importable, build_workbook fails loud
        so report()'s try/except can convert to a warning."""
        import sys

        monkeypatch.setitem(sys.modules, "openbb_techtrade", None)
        monkeypatch.setitem(sys.modules, "openbb_techtrade.reporting", None)

        from openbb_fmp_trading.reporting.xlsx_builder import (
            XLSXUnavailable, build_workbook,
        )
        with pytest.raises(XLSXUnavailable):
            build_workbook("s", events=[], metrics=None, output_path="ignored")


class TestPerSymbolPnLDecimal:
    """Review finding #3: PerSymbolPnL must use Decimal, not float."""

    def test_per_symbol_pnl_preserves_decimal_precision(self, tmp_path):
        pytest.importorskip("openpyxl")
        pytest.importorskip("openbb_techtrade")

        from openbb_fmp_trading.models.journal_events import FillEvent
        from openbb_fmp_trading.models.report import SessionMetrics
        from openbb_fmp_trading.reporting.xlsx_builder import build_workbook

        # Fills with a P&L that would lose precision under float summation
        events = [
            FillEvent(
                ts=datetime(2026, 7, 13, 14, 0, tzinfo=timezone.utc),
                session_id="s",
                payload={"symbol": "MSFT", "realized_pnl": "0.10"},
            ),
            FillEvent(
                ts=datetime(2026, 7, 13, 14, 5, tzinfo=timezone.utc),
                session_id="s",
                payload={"symbol": "MSFT", "realized_pnl": "0.20"},
            ),
        ]
        output = tmp_path / "end_of_day.xlsx"
        build_workbook(
            "s", events=events,
            metrics=SessionMetrics(realized_pnl=Decimal("0.30")),
            output_path=output,
        )

        from openpyxl import load_workbook
        wb = load_workbook(output)
        # PerSymbolPnL row for MSFT should be "0.30" under Decimal.
        # Under float summation it would be "0.30000000000000004"
        # (classic 0.1+0.2 IEEE754 artifact).
        pnl_row = wb["PerSymbolPnL"][2]  # header at row 1
        assert pnl_row[0].value == "MSFT"
        assert pnl_row[1].value == "0.30", (
            f"PerSymbolPnL P&L accumulation must use Decimal (P2). "
            f"Got {pnl_row[1].value!r} — float artifact suggests review "
            f"finding #3 wasn't fixed."
        )
