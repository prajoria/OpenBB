"""Typed story-artifact inputs and their deterministic renderings (#2217, N07).

Every renderer takes validated models and returns bytes; none touches the
filesystem except the TechTrade export, which writes into a private temporary
directory. Text that a spreadsheet would evaluate is quoted, HTML is escaped,
JSON is canonical and workbooks carry fixed timestamps, so equal inputs give
equal bytes and equal hashes.
"""

from __future__ import annotations

import csv
import html
import io
import json
import math
import re
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Annotated, Any, Literal

from pydantic import (
    BaseModel,
    BeforeValidator,
    ConfigDict,
    Field,
    field_validator,
    model_validator,
)

MAX_PREVIEW_COLUMNS = 50
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ArtifactKind = Literal["table", "chart", "report", "workbook"]

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
# Excel's own sheet-name rules: 1-31 characters, none of []:*?/\ .
_SHEET_NAME = re.compile(r"^[^\[\]:*?/\\]{1,31}$")
# A fixed zip and document timestamp keeps workbook bytes reproducible.
_ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)
_DOC_EPOCH = "2000-01-01T00:00:00Z"
_CORE_DATES = re.compile(r"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)")


def _json_safe(value: Any) -> Any:
    # A NaN in the input would make the 422 itself unencodable as JSON. Text
    # is a valid cell, so swap it for a mapping that no cell type accepts.
    if isinstance(value, float) and not math.isfinite(value):
        return {"non-finite number": str(value)}
    return value


Cell = Annotated[str | bool | int | float | None, BeforeValidator(_json_safe)]
Name = Annotated[
    str, Field(pattern=r"^[A-Za-z0-9][A-Za-z0-9 ._-]{0,79}$", max_length=80)
]
Text = Annotated[str, Field(max_length=4000)]


def _finite(value: Any) -> Any:
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("non-finite number")
    return value


class ArtifactRefOut(BaseModel):
    """The N03 ``ArtifactRef``: exactly its fields, nothing else."""

    model_config = ConfigDict(extra="forbid")

    run_id: str
    artifact_id: str
    media_type: str
    sha256: str
    size_bytes: int


class ArtifactRecord(BaseModel):
    """An artifact as its principal sees it; never a filesystem path."""

    model_config = ConfigDict(extra="forbid")

    ref: ArtifactRefOut
    name: str
    kind: ArtifactKind
    story_id: str
    stage_id: str | None
    created_at: datetime
    expires_at: datetime


class Table(BaseModel):
    """Columns and rows; every row matches the header."""

    model_config = ConfigDict(extra="forbid")

    columns: list[Annotated[str, Field(min_length=1, max_length=200)]] = Field(
        min_length=1, max_length=50
    )
    rows: list[list[Cell]] = Field(default_factory=list, max_length=5000)

    @field_validator("columns")
    @classmethod
    def _unique(cls, columns: list[str]) -> list[str]:
        if len(set(columns)) != len(columns):
            raise ValueError("column names must be unique")
        return columns

    @field_validator("rows")
    @classmethod
    def _bounded(cls, rows: list[list[Any]]) -> list[list[Any]]:
        for row in rows:
            for value in row:
                _finite(value)
                if isinstance(value, str) and len(value) > 2000:
                    raise ValueError("cell text is limited to 2000 characters")
        return rows

    @model_validator(mode="after")
    def _rectangular(self) -> Table:
        width = len(self.columns)
        for index, row in enumerate(self.rows):
            if len(row) != width:
                raise ValueError(f"row {index} has {len(row)} cells, expected {width}")
        return self


class Sheet(Table):
    """One workbook sheet."""

    name: str

    @field_validator("name")
    @classmethod
    def _excel_name(cls, name: str) -> str:
        if not _SHEET_NAME.fullmatch(name):
            raise ValueError("sheet names are 1-31 characters without []:*?/\\")
        return name


class WriteRequest(BaseModel):
    """Fields every artifact write shares: a label, its stage and retention."""

    model_config = ConfigDict(extra="forbid")

    name: Name
    stage_id: str | None = Field(default=None, max_length=200)
    retention_days: int = Field(default=30, ge=1, le=365)


class TableRequest(WriteRequest):
    """Render one table as CSV or canonical JSON."""

    format: Literal["csv", "json"] = "csv"
    table: Table


class WorkbookRequest(WriteRequest):
    """Render typed sheets as an XLSX workbook."""

    sheets: list[Sheet] = Field(min_length=1, max_length=10)

    @field_validator("sheets")
    @classmethod
    def _unique_sheets(cls, sheets: list[Sheet]) -> list[Sheet]:
        names = [s.name.lower() for s in sheets]
        if len(set(names)) != len(names):
            raise ValueError("sheet names must be unique")
        return sheets


class Series(BaseModel):
    """One named series of y values."""

    model_config = ConfigDict(extra="forbid")

    name: Annotated[str, Field(min_length=1, max_length=100)]
    values: list[Annotated[float | None, BeforeValidator(_json_safe)]] = Field(
        max_length=2000
    )

    @field_validator("values")
    @classmethod
    def _finite_values(cls, values: list[float | None]) -> list[float | None]:
        for value in values:
            _finite(value)
        return values


class Chart(BaseModel):
    """A chart as data plus a minimal spec, not an image."""

    model_config = ConfigDict(extra="forbid")

    kind: Literal["line", "bar", "scatter", "area"]
    title: str | None = Field(default=None, max_length=200)
    x_label: str | None = Field(default=None, max_length=100)
    y_label: str | None = Field(default=None, max_length=100)
    x: list[Annotated[str | float, BeforeValidator(_json_safe)]] = Field(
        min_length=1, max_length=2000
    )
    series: list[Series] = Field(min_length=1, max_length=10)

    @model_validator(mode="after")
    def _aligned(self) -> Chart:
        for value in self.x:
            _finite(value)
        for series in self.series:
            if len(series.values) != len(self.x):
                raise ValueError(
                    f"series '{series.name}' has {len(series.values)} values"
                    f" for {len(self.x)} x points"
                )
        return self


class ChartRequest(WriteRequest):
    """Store one chart spec with its data."""

    chart: Chart


class KeyValue(BaseModel):
    """One labelled value in a report section."""

    model_config = ConfigDict(extra="forbid")

    label: Annotated[str, Field(max_length=200)]
    value: Annotated[str, Field(max_length=1000)]


class Section(BaseModel):
    """One report section: text, facts, an optional table and citations."""

    model_config = ConfigDict(extra="forbid")

    heading: Annotated[str, Field(min_length=1, max_length=200)]
    paragraphs: list[Text] = Field(default_factory=list, max_length=20)
    key_values: list[KeyValue] = Field(default_factory=list, max_length=50)
    table: Table | None = None
    artifact_ids: list[str] = Field(default_factory=list, max_length=20)


class ReportRequest(WriteRequest):
    """Compose an escaped, self-contained HTML report."""

    title: Annotated[str, Field(min_length=1, max_length=200)]
    sections: list[Section] = Field(min_length=1, max_length=30)


class TechTradeWorkbookRequest(WriteRequest):
    """The six-sheet TechTrade export for these plans."""

    plans: list[dict[str, Any]] = Field(min_length=1, max_length=200)
    include_sheets: list[str] | None = Field(default=None, max_length=6)


class SheetView(BaseModel):
    """A bounded preview of one workbook sheet."""

    model_config = ConfigDict(extra="forbid")

    name: str
    max_row: int
    max_column: int
    charts: int
    rows: list[list[Any]]


def _neutralise(value: Any) -> Any:
    """Quote text a spreadsheet would otherwise evaluate as a formula."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _canonical_json(value: Any) -> bytes:
    return json.dumps(
        value, sort_keys=True, separators=(",", ":"), ensure_ascii=False
    ).encode("utf-8")


def render_csv(table: Table) -> bytes:
    """Render a table as CSV with formula-safe cells and LF line endings."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([_neutralise(c) for c in table.columns])
    for row in table.rows:
        writer.writerow([_neutralise(c) for c in row])
    return buffer.getvalue().encode("utf-8")


def render_table_json(table: Table) -> bytes:
    """Canonical JSON: sorted keys, compact separators."""
    return _canonical_json(table.model_dump(mode="json"))


def render_chart(chart: Chart) -> bytes:
    """Canonical JSON chart spec with its data."""
    return _canonical_json(
        {"schema": "openbb.story.chart/v1", **chart.model_dump(mode="json")}
    )


def _reproducible_xlsx(data: bytes) -> bytes:
    """Fix the zip entry times and document dates so equal inputs hash equal."""
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = _CORE_DATES.sub(
                    rf"\g<1>{_DOC_EPOCH}\g<3>", content.decode("utf-8")
                ).encode("utf-8")
            entry = zipfile.ZipInfo(info.filename, date_time=_ZIP_EPOCH)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            target.writestr(entry, content)
    return output.getvalue()


def render_workbook(sheets: list[Sheet]) -> bytes:
    """XLSX with a bold header row per sheet and formula-safe cells."""
    from openpyxl import Workbook  # pylint: disable=import-outside-toplevel
    from openpyxl.styles import Font  # pylint: disable=import-outside-toplevel

    workbook = Workbook()
    workbook.remove(workbook.active)
    bold = Font(bold=True)
    for sheet in sheets:
        worksheet = workbook.create_sheet(title=sheet.name)
        worksheet.append([_neutralise(c) for c in sheet.columns])
        for cell in worksheet[1]:
            cell.font = bold
        for row in sheet.rows:
            worksheet.append([_neutralise(c) for c in row])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return _reproducible_xlsx(buffer.getvalue())


def render_techtrade_workbook(
    plans: list[dict[str, Any]], include_sheets: list[str] | None
) -> bytes:
    """Render the TechTrade six-sheet export with formula-safe text cells."""
    # pylint: disable=import-outside-toplevel
    from openbb_techtrade.models import ExportConfig, TradePlan
    from openbb_techtrade.reporting import excel_export

    validated = [TradePlan.model_validate(plan) for plan in plans]
    config = (
        ExportConfig(include_sheets=include_sheets)
        if include_sheets is not None
        else ExportConfig()
    )
    unknown = set(config.include_sheets) - set(excel_export.CANONICAL_SHEETS)
    if unknown or not config.include_sheets:
        raise ValueError(
            f"include_sheets must be drawn from {list(excel_export.CANONICAL_SHEETS)}"
        )
    sheets = excel_export._resolved_sheets(config)  # pylint: disable=protected-access
    ordered = excel_export._sorted_plans(validated)  # pylint: disable=protected-access
    frames = excel_export._build_dataframes(  # pylint: disable=protected-access
        ordered, None, sheets
    )
    for name, frame in frames.items():
        frames[name] = frame.map(_neutralise)
    as_of = ordered[0].as_of if ordered else None
    with tempfile.TemporaryDirectory() as scratch:
        path = Path(scratch) / "workbook.xlsx"
        excel_export._write_workbook_openpyxl(  # pylint: disable=protected-access
            path, sheets, frames, config, as_of
        )
        data = path.read_bytes()
    return _reproducible_xlsx(data)


_REPORT_STYLE = (
    "body{font-family:system-ui,sans-serif;max-width:60rem;margin:2rem auto;"
    "padding:0 1rem;line-height:1.5}table{border-collapse:collapse}"
    "th,td{border:1px solid #ccc;padding:.25rem .5rem;text-align:left}"
    "code{font-size:.85em}"
)


def _html_table(table: Table) -> str:
    head = "".join(f"<th>{html.escape(c)}</th>" for c in table.columns)
    body = "".join(
        "<tr>"
        + "".join(f"<td>{html.escape('' if c is None else str(c))}</td>" for c in row)
        + "</tr>"
        for row in table.rows
    )
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def render_report(request: ReportRequest, cited: dict[str, ArtifactRecord]) -> bytes:
    """Self-contained HTML; every caller string is escaped."""
    parts = [
        '<!doctype html><html lang="en"><head><meta charset="utf-8">',
        '<meta http-equiv="Content-Security-Policy"'
        " content=\"default-src 'none'; style-src 'unsafe-inline'\">",
        f"<title>{html.escape(request.title)}</title>",
        f"<style>{_REPORT_STYLE}</style></head><body>",
        f"<h1>{html.escape(request.title)}</h1>",
    ]
    for section in request.sections:
        parts.append(f"<section><h2>{html.escape(section.heading)}</h2>")
        parts.extend(f"<p>{html.escape(p)}</p>" for p in section.paragraphs)
        if section.key_values:
            parts.append("<dl>")
            for item in section.key_values:
                parts.append(
                    f"<dt>{html.escape(item.label)}</dt>"
                    f"<dd>{html.escape(item.value)}</dd>"
                )
            parts.append("</dl>")
        if section.table is not None:
            parts.append(_html_table(section.table))
        if section.artifact_ids:
            parts.append('<ul class="artifacts">')
            for artifact_id in section.artifact_ids:
                record = cited[artifact_id]
                parts.append(
                    f"<li>{html.escape(record.name)} ({html.escape(record.kind)},"
                    f" {html.escape(record.ref.media_type)})"
                    f" <code>{record.ref.artifact_id}</code>"
                    f" sha256 <code>{record.ref.sha256}</code></li>"
                )
            parts.append("</ul>")
        parts.append("</section>")
    parts.append("</body></html>")
    return "".join(parts).encode("utf-8")


def _cell_out(value: Any) -> Any:
    if value is None or isinstance(value, bool | int | float | str):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def workbook_view(data: bytes, max_rows: int) -> list[SheetView]:
    """Bounded preview of a workbook this store wrote and hash-verified."""
    from openpyxl import load_workbook  # pylint: disable=import-outside-toplevel

    workbook = load_workbook(io.BytesIO(data))
    views = []
    for worksheet in workbook.worksheets:
        # iter_rows pads to the requested bounds, so clamp to the sheet's size.
        rows = [
            [_cell_out(v) for v in row]
            for row in worksheet.iter_rows(
                max_row=min(max_rows, worksheet.max_row),
                max_col=min(MAX_PREVIEW_COLUMNS, worksheet.max_column),
                values_only=True,
            )
        ]
        views.append(
            SheetView(
                name=worksheet.title,
                max_row=worksheet.max_row,
                max_column=worksheet.max_column,
                charts=len(worksheet._charts),  # pylint: disable=protected-access
                rows=rows,
            )
        )
    return views
