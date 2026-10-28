"""Bounded, run-scoped story artifacts (#2217, N07)."""

# pylint: disable=redefined-outer-name,protected-access

import hashlib
import importlib.util
import io
import json
import os
import sys
import zipfile
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openbb_mcp_server.story_adapters import artifacts
from openbb_mcp_server.story_adapters.artifacts import (
    ArtifactStore,
    InvalidArtifactId,
    compose_story_artifacts_app,
)
from openbb_mcp_server.story_adapters.registry import compose_story_registry_app
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[5]
MANIFESTS = (
    REPO_ROOT
    / "openbb_platform"
    / "extensions"
    / "agents"
    / "openbb_agents"
    / "notebook_guide"
    / "stories"
)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _n03_contracts():
    # The N03 contracts depend only on pydantic; load them by path so this
    # suite checks the real schema even where openbb-agents is not installed.
    spec = importlib.util.spec_from_file_location(
        "_n03_contracts", MANIFESTS.parent / "contracts.py"
    )
    module = importlib.util.module_from_spec(spec)
    # Pydantic resolves the module's postponed annotations through sys.modules.
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


ArtifactRef = _n03_contracts().ArtifactRef


def _client(principal: str, store: Path, root: Path) -> TestClient:
    app = compose_story_registry_app(
        FastAPI(), principal=principal, manifest_dir=MANIFESTS, store_path=store
    )
    app = compose_story_artifacts_app(
        app, principal=principal, store_path=store, artifact_root=root
    )
    return TestClient(app)


@pytest.fixture
def harness(tmp_path):
    """Two principals sharing one run store and one artifact root."""
    store = tmp_path / "runs.sqlite"
    root = tmp_path / "artifacts"
    alice = _client("alice", store, root)
    bob = _client("bob", store, root)

    def run(story_id="techtrade/07", client=alice):
        response = client.post("/api/v1/story/runs", json={"story_id": story_id})
        assert response.status_code == 200, response.text
        return response.json()["run_id"]

    def post(run_id, kind, body, client=alice):
        return client.post(f"/api/v1/story/runs/{run_id}/artifacts/{kind}", json=body)

    def read(run_id, artifact_id, client=alice, **params):
        return client.get(
            f"/api/v1/story/runs/{run_id}/artifacts/{artifact_id}", params=params
        )

    def table(run_id, name="positions", fmt="csv", client=alice, **extra):
        body = {
            "name": name,
            "format": fmt,
            "table": {"columns": ["symbol", "weight"], "rows": [["AAA", 0.5]]},
            **extra,
        }
        response = post(run_id, "table", body, client=client)
        assert response.status_code == 200, response.text
        return response.json()

    def raw(artifact_id):
        matches = list(root.rglob(f"{artifact_id}.*"))
        assert len(matches) == 1
        return matches[0]

    return SimpleNamespace(
        alice=alice,
        bob=bob,
        root=root,
        store=store,
        run=run,
        post=post,
        read=read,
        table=table,
        raw=raw,
    )


@pytest.fixture
def artifact_harness(tmp_path):
    """Bind the bare store to one principal and run."""
    store = ArtifactStore(tmp_path / "artifacts", tmp_path / "runs.sqlite")
    run_id = "a" * 32
    return SimpleNamespace(
        InvalidArtifactId=InvalidArtifactId,
        read=lambda artifact_id: store.read("alice", run_id, artifact_id),
    )


def test_artifact_id_cannot_be_a_path(artifact_harness):
    """The plan's anchor case: an artifact ID is never a filesystem path."""
    with pytest.raises(artifact_harness.InvalidArtifactId):
        artifact_harness.read("../outside.xlsx")


@pytest.mark.parametrize(
    "artifact_id", ["..%5Coutside", "outside.xlsx", "A" * 32, "a" * 31, "a" * 33]
)
def test_route_rejects_artifact_ids_that_are_not_server_issued(harness, artifact_id):
    """Only the 32-hex IDs the server issues are accepted; nothing is opened."""
    run_id = harness.run()
    response = harness.read(run_id, artifact_id)
    assert response.status_code == 422, response.text


def test_table_csv_ref_matches_the_n03_contract_and_the_bytes(harness):
    """The returned ref validates as N03's ArtifactRef and hashes the file."""
    run_id = harness.run()
    record = harness.table(run_id)
    ref = ArtifactRef.model_validate(record["ref"])
    data = harness.raw(ref.artifact_id).read_bytes()
    assert data == b"symbol,weight\nAAA,0.5\n"
    assert ref.run_id == run_id
    assert ref.media_type == "text/csv"
    assert ref.size_bytes == len(data)
    assert ref.sha256 == hashlib.sha256(data).hexdigest()
    assert record["kind"] == "table"
    assert record["story_id"] == "techtrade/07"
    read = harness.read(run_id, ref.artifact_id).json()
    assert read["text"] == "symbol,weight\nAAA,0.5\n"
    assert read["truncated"] is False
    assert read["workbook"] is None


def test_table_json_is_canonical(harness):
    """JSON tables are written sorted and compact, so the hash is stable."""
    run_id = harness.run()
    record = harness.table(run_id, fmt="json")
    assert record["ref"]["media_type"] == "application/json"
    data = harness.raw(record["ref"]["artifact_id"]).read_bytes()
    assert data == b'{"columns":["symbol","weight"],"rows":[["AAA",0.5]]}'


INJECTIONS = ["=1+1", "+SUM(A1)", "-2+3", "@cmd", "\tx", "\r=y"]


def test_formula_injection_is_neutralised_in_csv(harness):
    """Leading formula characters are quoted; real numbers stay numbers."""
    run_id = harness.run()
    body = {
        "name": "inj",
        "format": "csv",
        "table": {"columns": ["v"], "rows": [[v] for v in INJECTIONS] + [[-5]]},
    }
    record = harness.post(run_id, "table", body).json()
    lines = harness.raw(record["ref"]["artifact_id"]).read_bytes().decode()
    assert "'=1+1" in lines
    assert "'+SUM(A1)" in lines
    assert "'-2+3" in lines
    assert "'@cmd" in lines
    assert "'\tx" in lines
    assert lines.rstrip("\n").endswith("\n-5")


def test_formula_injection_is_neutralised_in_xlsx(harness):
    """No workbook cell written from caller input is ever a formula."""
    run_id = harness.run()
    body = {
        "name": "inj",
        "sheets": [
            {
                "name": "Data",
                "columns": ["=HYPERLINK(1)", "n"],
                "rows": [[v, -5] for v in INJECTIONS],
            }
        ],
    }
    response = harness.post(run_id, "workbook", body)
    assert response.status_code == 200, response.text
    path = harness.raw(response.json()["ref"]["artifact_id"])
    sheet = load_workbook(path)["Data"]
    cells = [cell for row in sheet.iter_rows() for cell in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    assert sheet["A1"].value == "'=HYPERLINK(1)"
    assert sheet["A2"].value == "'=1+1"
    assert sheet["B2"].value == -5


def test_workbook_sheet_and_cell_contract_is_exact(harness):
    """Sheets keep their order, a bold header row, and typed cell values."""
    run_id = harness.run()
    body = {
        "name": "basket",
        "sheets": [
            {
                "name": "Orders",
                "columns": ["symbol", "qty", "limit", "filled"],
                "rows": [["AAA", 10, 12.5, True], ["BBB", 3, None, False]],
            },
            {"name": "Summary", "columns": ["metric", "value"], "rows": [["n", 2]]},
        ],
    }
    record = harness.post(run_id, "workbook", body).json()
    assert record["ref"]["media_type"] == XLSX
    assert record["kind"] == "workbook"
    workbook = load_workbook(harness.raw(record["ref"]["artifact_id"]))
    assert workbook.sheetnames == ["Orders", "Summary"]
    orders = workbook["Orders"]
    assert [c.value for c in orders[1]] == ["symbol", "qty", "limit", "filled"]
    assert all(c.font.bold for c in orders[1])
    assert [c.value for c in orders[2]] == ["AAA", 10, 12.5, True]
    assert [c.value for c in orders[3]] == ["BBB", 3, None, False]
    assert orders.max_row == 3
    view = harness.read(run_id, record["ref"]["artifact_id"]).json()["workbook"]
    assert [s["name"] for s in view] == ["Orders", "Summary"]
    assert view[0]["rows"] == [
        ["symbol", "qty", "limit", "filled"],
        ["AAA", 10, 12.5, True],
        ["BBB", 3, None, False],
    ]
    assert view[0]["max_row"] == 3
    assert view[0]["max_column"] == 4
    assert view[0]["charts"] == 0


def test_workbook_bytes_are_deterministic(harness):
    """Identical inputs give identical bytes, whatever the wall clock says."""
    run_id = harness.run()
    sheet = {"name": "S", "columns": ["a"], "rows": [[1]]}
    first = harness.post(run_id, "workbook", {"name": "one", "sheets": [sheet]})
    second = harness.post(run_id, "workbook", {"name": "two", "sheets": [sheet]})
    assert first.json()["ref"]["sha256"] == second.json()["ref"]["sha256"]
    _assert_reproducible(harness.raw(first.json()["ref"]["artifact_id"]))


def _assert_reproducible(path):
    # Two writes in the same second would hash equal anyway, so check the
    # clock-dependent fields directly: zip entry times and document dates.
    with zipfile.ZipFile(path) as archive:
        assert {i.date_time for i in archive.infolist()} == {(1980, 1, 1, 0, 0, 0)}
        core = archive.read("docProps/core.xml").decode()
    assert core.count("2000-01-01T00:00:00Z") == 2


def test_workbook_read_is_bounded_by_max_rows(harness):
    """A read previews at most ``max_rows`` rows but reports the real size."""
    run_id = harness.run()
    rows = [[i] for i in range(30)]
    body = {"name": "big", "sheets": [{"name": "S", "columns": ["i"], "rows": rows}]}
    artifact_id = harness.post(run_id, "workbook", body).json()["ref"]["artifact_id"]
    view = harness.read(run_id, artifact_id, max_rows=5).json()["workbook"][0]
    assert len(view["rows"]) == 5
    assert view["max_row"] == 31
    assert harness.read(run_id, artifact_id, max_rows=201).status_code == 422


@pytest.mark.parametrize("bad", ["Sheet[1]", "a/b", "x" * 32, ""])
def test_workbook_rejects_sheet_names_excel_cannot_hold(harness, bad):
    """Sheet names follow Excel's rules instead of being silently renamed."""
    run_id = harness.run()
    body = {"name": "w", "sheets": [{"name": bad, "columns": ["a"], "rows": []}]}
    assert harness.post(run_id, "workbook", body).status_code == 422


def test_ragged_rows_and_duplicate_columns_are_rejected(harness):
    """A row must match the header; header names must be unique."""
    run_id = harness.run()
    ragged = {"columns": ["a", "b"], "rows": [["x"]]}
    dup = {"columns": ["a", "a"], "rows": []}
    for table in (ragged, dup):
        body = {"name": "t", "format": "csv", "table": table}
        assert harness.post(run_id, "table", body).status_code == 422


def test_chart_is_a_canonical_spec_with_its_data(harness):
    """Charts are stored as deterministic JSON specs, not rendered images."""
    run_id = harness.run()
    chart = {
        "kind": "line",
        "title": "Equity",
        "x_label": "date",
        "y_label": "value",
        "x": ["2024-01-02", "2024-01-03"],
        "series": [{"name": "basket", "values": [100.0, 101.5]}],
    }
    response = harness.post(run_id, "chart", {"name": "equity", "chart": chart})
    assert response.status_code == 200, response.text
    record = response.json()
    assert record["kind"] == "chart"
    data = harness.raw(record["ref"]["artifact_id"]).read_bytes()
    assert json.loads(data) == {"schema": "openbb.story.chart/v1", **chart}
    assert (
        data
        == json.dumps(
            {"schema": "openbb.story.chart/v1", **chart},
            sort_keys=True,
            separators=(",", ":"),
        ).encode()
    )


def test_chart_series_must_match_the_x_axis(harness):
    """A series shorter or longer than x is rejected, not padded."""
    run_id = harness.run()
    chart = {"kind": "bar", "x": ["a", "b"], "series": [{"name": "s", "values": [1]}]}
    assert (
        harness.post(run_id, "chart", {"name": "c", "chart": chart}).status_code == 422
    )


@pytest.mark.parametrize("literal", ["NaN", "Infinity", "-Infinity"])
def test_non_finite_numbers_are_a_clean_422(harness, literal):
    """NaN and infinity are rejected with a JSON-encodable error."""
    run_id = harness.run()
    table = '{"columns":["a"],"rows":[[' + literal + "]]}"
    body = '{"name":"t","format":"csv","table":' + table + "}"
    response = harness.alice.post(
        f"/api/v1/story/runs/{run_id}/artifacts/table",
        content=body,
        headers={"content-type": "application/json"},
    )
    assert response.status_code == 422
    assert "non-finite" in response.text


def test_report_escapes_text_and_cites_same_run_artifacts(harness):
    """Reports are escaped HTML that cite artifacts by ID and hash only."""
    run_id = harness.run()
    table = harness.table(run_id)
    body = {
        "name": "report",
        "title": "Monday <routine>",
        "sections": [
            {
                "heading": "Basket",
                "paragraphs": ["<script>alert(1)</script>"],
                "key_values": [{"label": "Names", "value": "10"}],
                "table": {"columns": ["a"], "rows": [["<b>x</b>"]]},
                "artifact_ids": [table["ref"]["artifact_id"]],
            }
        ],
    }
    response = harness.post(run_id, "report", body)
    assert response.status_code == 200, response.text
    record = response.json()
    assert record["ref"]["media_type"] == "text/html"
    html = harness.raw(record["ref"]["artifact_id"]).read_text(encoding="utf-8")
    assert "<script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "Monday &lt;routine&gt;" in html
    assert "&lt;b&gt;x&lt;/b&gt;" in html
    assert table["ref"]["sha256"] in html
    assert "positions" in html
    assert "Content-Security-Policy" in html


def test_report_cannot_cite_another_runs_artifact(harness):
    """A citation must name an artifact of the same run."""
    other = harness.table(harness.run(), name="other")
    run_id = harness.run()
    body = {
        "name": "report",
        "title": "t",
        "sections": [{"heading": "h", "artifact_ids": [other["ref"]["artifact_id"]]}],
    }
    response = harness.post(run_id, "report", body)
    assert response.status_code == 422
    assert not list(harness.root.rglob("*.html"))


def test_cross_run_and_cross_principal_reads_are_denied(harness):
    """An artifact is reachable only through its own run and principal."""
    run_a = harness.run()
    run_b = harness.run()
    artifact_id = harness.table(run_a)["ref"]["artifact_id"]
    assert harness.read(run_b, artifact_id).status_code == 404
    assert harness.read(run_a, artifact_id, client=harness.bob).status_code == 404
    listed = harness.bob.get("/api/v1/story/artifacts").json()
    assert listed == []
    assert harness.read(run_a, artifact_id).status_code == 200


def test_writes_need_a_run_the_principal_owns(harness):
    """Bob cannot write into Alice's run, and unknown runs are 404."""
    run_id = harness.run()
    body = {"name": "t", "format": "csv", "table": {"columns": ["a"], "rows": []}}
    assert harness.post(run_id, "table", body, client=harness.bob).status_code == 404
    assert harness.post("f" * 32, "table", body).status_code == 404
    assert not list(harness.root.rglob("*.csv"))


def test_stage_must_belong_to_the_run(harness):
    """An artifact may name the stage that produced it, if the run has it."""
    run_id = harness.run()
    ok = harness.table(run_id, stage_id="6-the-6-sheet-xlsx-workbook")
    assert ok["stage_id"] == "6-the-6-sheet-xlsx-workbook"
    body = {
        "name": "t2",
        "format": "csv",
        "table": {"columns": ["a"], "rows": []},
        "stage_id": "no-such-stage",
    }
    assert harness.post(run_id, "table", body).status_code == 422


def test_names_are_unique_per_run_so_nothing_is_overwritten(harness):
    """A second artifact with the same name is a 409 and the first survives."""
    run_id = harness.run()
    first = harness.table(run_id)
    body = {
        "name": "positions",
        "format": "csv",
        "table": {"columns": ["x"], "rows": [["changed"]]},
    }
    assert harness.post(run_id, "table", body).status_code == 409
    data = harness.raw(first["ref"]["artifact_id"]).read_bytes()
    assert hashlib.sha256(data).hexdigest() == first["ref"]["sha256"]
    assert len(list(harness.root.rglob("*.csv"))) == 1


@pytest.mark.parametrize("bad", ["../x", "a/b", "a\\b", " lead", "", "x" * 81])
def test_names_are_labels_not_paths(harness, bad):
    """A display name never reaches the filesystem and is validated."""
    run_id = harness.run()
    body = {"name": bad, "format": "csv", "table": {"columns": ["a"], "rows": []}}
    assert harness.post(run_id, "table", body).status_code == 422


def test_oversized_artifacts_are_rejected_before_writing(harness, monkeypatch):
    """Rendered bytes above the limit are a 413 and leave no file behind."""
    monkeypatch.setattr(artifacts, "MAX_ARTIFACT_BYTES", 16)
    run_id = harness.run()
    body = {
        "name": "big",
        "format": "csv",
        "table": {"columns": ["a"], "rows": [["x" * 40]]},
    }
    assert harness.post(run_id, "table", body).status_code == 413
    assert not [p for p in harness.root.rglob("*") if p.is_file()]


def test_runs_have_an_artifact_quota(harness, monkeypatch):
    """A run holds a bounded number of artifacts."""
    monkeypatch.setattr(artifacts, "MAX_ARTIFACTS_PER_RUN", 2)
    run_id = harness.run()
    harness.table(run_id, name="one")
    harness.table(run_id, name="two")
    body = {"name": "three", "format": "csv", "table": {"columns": ["a"], "rows": []}}
    assert harness.post(run_id, "table", body).status_code == 413


def test_runs_have_a_byte_quota(harness, monkeypatch):
    """A run's artifacts are bounded in total size, not just count."""
    monkeypatch.setattr(artifacts, "MAX_RUN_BYTES", 40)
    run_id = harness.run()
    harness.table(run_id, name="one")
    body = {
        "name": "two",
        "format": "csv",
        "table": {"columns": ["symbol", "weight"], "rows": [["BBBBBBBB", 0.25]]},
    }
    assert harness.post(run_id, "table", body).status_code == 413


def _link_dir(link: Path, target: Path) -> None:
    try:
        os.symlink(target, link, target_is_directory=True)
        return
    except (OSError, NotImplementedError):
        pass
    if sys.platform == "win32":
        import _winapi  # pylint: disable=import-outside-toplevel

        _winapi.CreateJunction(str(target), str(link))
        return
    pytest.skip("cannot create a directory link here")


def test_a_linked_run_directory_cannot_redirect_writes(harness, tmp_path):
    """A planted symlink or junction for the run directory is refused."""
    run_id = harness.run()
    outside = tmp_path / "outside"
    outside.mkdir()
    principal_dir = harness.root / ArtifactStore.principal_key("alice")
    principal_dir.mkdir(parents=True)
    _link_dir(principal_dir / run_id, outside)
    body = {"name": "t", "format": "csv", "table": {"columns": ["a"], "rows": []}}
    response = harness.post(run_id, "table", body)
    assert response.status_code == 409, response.text
    assert not list(outside.iterdir())


def test_a_linked_store_root_is_refused(tmp_path):
    """The artifact root itself may not be a link to somewhere else."""
    outside = tmp_path / "outside"
    outside.mkdir()
    root = tmp_path / "root"
    _link_dir(root, outside)
    client = _client("alice", tmp_path / "runs.sqlite", root)
    run_id = client.post("/api/v1/story/runs", json={"story_id": "techtrade/07"})
    body = {"name": "t", "format": "csv", "table": {"columns": ["a"], "rows": []}}
    response = client.post(
        f"/api/v1/story/runs/{run_id.json()['run_id']}/artifacts/table", json=body
    )
    assert response.status_code == 409, response.text
    assert not list(outside.iterdir())


def test_a_file_changed_on_disk_is_not_served(harness):
    """Reads re-hash the bytes; tampering is a 409, never served content."""
    run_id = harness.run()
    artifact_id = harness.table(run_id)["ref"]["artifact_id"]
    path = harness.raw(artifact_id)
    tampered = b"symbol,weight\nZZZ,0.9\n"
    # Same length as the original, so only the hash can tell them apart.
    assert len(tampered) == path.stat().st_size
    path.write_bytes(tampered)
    response = harness.read(run_id, artifact_id)
    assert response.status_code == 409
    assert "ZZZ" not in response.text


def test_create_exclusive_never_replaces_a_file(tmp_path):
    """Publishing a file fails, and leaves both files, if the name is taken."""
    final = tmp_path / "x.csv"
    final.write_bytes(b"original")
    with pytest.raises(artifacts.ArtifactConflict):
        artifacts._create_exclusive(tmp_path, final, b"replacement")
    assert final.read_bytes() == b"original"
    assert [p.name for p in tmp_path.iterdir()] == ["x.csv"]


def test_expired_artifacts_are_gone_and_purged(harness, monkeypatch):
    """Past retention an artifact reads as 410, leaves listings, then its file goes."""
    run_id = harness.run()
    artifact_id = harness.table(run_id, retention_days=1)["ref"]["artifact_id"]
    path = harness.raw(artifact_id)
    real_now = artifacts._now
    monkeypatch.setattr(artifacts, "_now", lambda: real_now() + timedelta(days=2))
    assert harness.read(run_id, artifact_id).status_code == 410
    listed = harness.alice.get("/api/v1/story/artifacts", params={"run_id": run_id})
    assert listed.json() == []
    harness.table(run_id, name="later")
    assert not path.exists()


def test_retention_is_bounded(harness):
    """Retention is between one day and a year."""
    run_id = harness.run()
    for days in (0, 366):
        body = {
            "name": f"r{days}",
            "format": "csv",
            "table": {"columns": ["a"], "rows": []},
            "retention_days": days,
        }
        assert harness.post(run_id, "table", body).status_code == 422
    record = harness.table(run_id)
    assert record["expires_at"] > record["created_at"]


def test_listing_is_scoped_by_run_and_story(harness):
    """List one run, or every run of the principal filtered by story or kind."""
    run_a = harness.run("techtrade/02")
    run_b = harness.run("techtrade/06")
    harness.table(run_a, name="a")
    chart = {"kind": "bar", "x": ["a"], "series": [{"name": "s", "values": [1]}]}
    harness.post(run_b, "chart", {"name": "b", "chart": chart})
    by_run = harness.alice.get("/api/v1/story/artifacts", params={"run_id": run_a})
    assert [r["name"] for r in by_run.json()] == ["a"]
    by_story = harness.alice.get(
        "/api/v1/story/artifacts", params={"story_id": "techtrade/02"}
    )
    assert [r["ref"]["run_id"] for r in by_story.json()] == [run_a]
    by_kind = harness.alice.get("/api/v1/story/artifacts", params={"kind": "chart"})
    assert [r["name"] for r in by_kind.json()] == ["b"]
    everything = harness.alice.get("/api/v1/story/artifacts").json()
    assert {r["name"] for r in everything} == {"a", "b"}
    unknown = harness.alice.get("/api/v1/story/artifacts", params={"run_id": "f" * 32})
    assert unknown.status_code == 404
    other = harness.alice.get("/api/v1/story/artifacts", params={"run_id": run_a})
    assert (
        harness.bob.get("/api/v1/story/artifacts", params={"run_id": run_a}).status_code
        == 404
    )
    assert other.status_code == 200


def test_responses_never_carry_a_filesystem_path(harness):
    """No response names the store root or any local path."""
    run_id = harness.run()
    record = harness.table(run_id)
    read = harness.read(run_id, record["ref"]["artifact_id"])
    listed = harness.alice.get("/api/v1/story/artifacts")
    blob = json.dumps([record, read.json(), listed.json()])
    for marker in (str(harness.root), harness.root.as_posix(), harness.root.name):
        assert marker not in blob


def test_composing_the_routes_writes_nothing(tmp_path):
    """The store is lazy: composing the routes has no filesystem side effect."""
    root = tmp_path / "artifacts"
    compose_story_artifacts_app(
        FastAPI(),
        principal="alice",
        store_path=tmp_path / "r.sqlite",
        artifact_root=root,
    )
    assert not root.exists()
    assert not (tmp_path / "r.sqlite").exists()


def test_the_artifact_routes_are_the_seven_mcp_tools(tmp_path):
    """Every route is an authenticated MCP tool with a stable name."""
    app = compose_story_artifacts_app(
        FastAPI(), principal="alice", store_path=tmp_path / "r.sqlite"
    )
    names = {
        route.openapi_extra["mcp_config"]["name"]
        for route in app.routes
        if getattr(route, "openapi_extra", None)
    }
    assert names == {
        "story_artifact_table",
        "story_artifact_chart",
        "story_artifact_report",
        "story_artifact_workbook",
        "story_artifact_techtrade_workbook",
        "story_artifact_list",
        "story_artifact_read",
    }
    assert all(
        route.dependencies
        for route in app.routes
        if getattr(route, "openapi_extra", None)
    )


def test_an_empty_principal_is_refused():
    """Artifacts are always owned by an authenticated principal."""
    with pytest.raises(ValueError):
        compose_story_artifacts_app(FastAPI(), principal="")


# --- TechTrade workbook: the existing six-sheet export, inside the run ------


def _flat_plan(symbol="MSFT", caveats="None."):
    from openbb_techtrade.models import (  # pylint: disable=import-outside-toplevel
        EntryExitRule,
        IndicatorVote,
        MoverSignal,
        Recommendation,
        TradePlan,
    )

    as_of = date(2024, 1, 12)
    entry = Decimal("121.40")
    signal = MoverSignal(
        symbol=symbol,
        segment="Information Technology",
        as_of=as_of,
        score=0.1,
        direction="flat",
        votes=[IndicatorVote(family="trend", name="macd_hist", vote=0.1, weight=0.4)],
        rank_in_segment=1,
    )
    recommendation = Recommendation(
        symbol=symbol,
        segment="Information Technology",
        as_of=as_of,
        action="HOLD/FLAT",
        conviction="Low",
        score=0.1,
        entry_price=entry,
        stop_price=entry,
        target_price=entry,
        stop_distance_pct=0.0,
        target_distance_pct=0.0,
        risk_reward=0.0,
        atr=1.9,
        position_size=Decimal(0),
        risk_per_share=Decimal(0),
        risk_pct_of_notional=0.0,
        time_stop_bars=None,
        reasoning=f"Hold {symbol}.",
        top_factors=[],
        caveats=caveats,
    )
    plan = TradePlan(
        symbol=symbol,
        segment="Information Technology",
        as_of=as_of,
        signal=signal,
        rule=EntryExitRule(),
        position_size=Decimal(0),
        orders=[],
        simulated_fills=[],
        recommendation=recommendation,
    )
    return json.loads(plan.model_dump_json())


def test_techtrade_workbook_keeps_the_export_layout(harness):
    """The run's workbook is the six-sheet TechTrade export, title row and all."""
    from openbb_techtrade.reporting.excel_export import (  # pylint: disable=import-outside-toplevel
        CANONICAL_SHEETS,
        DISCLAIMER,
    )

    run_id = harness.run("techtrade/02")
    body = {"name": "morning", "plans": [_flat_plan()]}
    response = harness.post(run_id, "techtrade_workbook", body)
    assert response.status_code == 200, response.text
    record = response.json()
    assert record["kind"] == "workbook"
    workbook = load_workbook(harness.raw(record["ref"]["artifact_id"]))
    assert tuple(workbook.sheetnames) == CANONICAL_SHEETS
    recs = workbook["Recommendations"]
    assert (
        recs["A1"].value == "OpenBB TechTrade - Top-Mover Recommendations - 2024-01-12"
    )
    assert recs["A2"].value == DISCLAIMER
    header = [c.value for c in recs[4]]
    symbol_col = header.index("Symbol") + 1
    assert recs.cell(row=5, column=symbol_col).value == "MSFT"
    view = harness.read(run_id, record["ref"]["artifact_id"]).json()["workbook"]
    assert [s["name"] for s in view] == list(CANONICAL_SHEETS)
    _assert_reproducible(harness.raw(record["ref"]["artifact_id"]))


def test_techtrade_workbook_neutralises_formulas_in_plan_text(harness):
    """Agent-supplied plan text cannot become a live formula in the export."""
    run_id = harness.run("techtrade/02")
    plan = _flat_plan(caveats='=HYPERLINK("http://example.invalid","x")')
    record = harness.post(run_id, "techtrade_workbook", {"name": "m", "plans": [plan]})
    assert record.status_code == 200, record.text
    workbook = load_workbook(harness.raw(record.json()["ref"]["artifact_id"]))
    formulas = [
        f"{ws.title}!{c.coordinate}"
        for ws in workbook.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.data_type == "f"
    ]
    assert formulas == []
    values = [
        c.value for ws in workbook.worksheets for row in ws.iter_rows() for c in row
    ]
    assert '\'=HYPERLINK("http://example.invalid","x")' in values


def test_techtrade_workbook_subsets_sheets_in_canonical_order(harness):
    """include_sheets subsets the export but never reorders it."""
    run_id = harness.run("techtrade/02")
    body = {
        "name": "m",
        "plans": [_flat_plan()],
        "include_sheets": ["Summary", "Recommendations"],
    }
    record = harness.post(run_id, "techtrade_workbook", body).json()
    workbook = load_workbook(harness.raw(record["ref"]["artifact_id"]))
    assert workbook.sheetnames == ["Recommendations", "Summary"]
    body["include_sheets"] = ["Nope"]
    body["name"] = "m2"
    assert harness.post(run_id, "techtrade_workbook", body).status_code == 422


def test_techtrade_workbook_rejects_malformed_plans(harness):
    """Plans are validated as TradePlan; a bad one is a 422 and writes nothing."""
    run_id = harness.run("techtrade/02")
    bad = _flat_plan()
    del bad["recommendation"]
    response = harness.post(run_id, "techtrade_workbook", {"name": "m", "plans": [bad]})
    assert response.status_code == 422
    assert (
        harness.post(
            run_id, "techtrade_workbook", {"name": "m", "plans": []}
        ).status_code
        == 422
    )
    assert not list(harness.root.rglob("*.xlsx"))


def test_store_reads_back_with_bytes_for_workbooks(tmp_path):
    """The store returns the bytes it hashed, so callers never touch paths."""
    store = ArtifactStore(tmp_path / "a", tmp_path / "r.sqlite")
    run_id = "b" * 32
    record = store.write(
        "alice",
        run_id,
        story_id="techtrade/07",
        name="w",
        kind="table",
        media_type="text/csv",
        extension="csv",
        payload=b"a\n1\n",
        retention_days=30,
    )
    got, data = store.read("alice", run_id, record.ref.artifact_id)
    assert data == b"a\n1\n"
    assert got == record
    assert io.BytesIO(data).read() == b"a\n1\n"
    # The store scopes by principal itself, not only behind the routes.
    with pytest.raises(artifacts.ArtifactNotFound):
        store.read("bob", run_id, record.ref.artifact_id)
    assert store.list("bob") == []
