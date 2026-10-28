"""TechTrade order batches and paper-sink artifacts (#2332, N09b)."""

# pylint: disable=redefined-outer-name,protected-access,import-outside-toplevel

import io
import json
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openbb_mcp_server.story_adapters import paper_batch
from openbb_mcp_server.story_adapters.artifacts import compose_story_artifacts_app
from openbb_mcp_server.story_adapters.paper_batch import compose_story_paper_batch_app
from openbb_mcp_server.story_adapters.registry import compose_story_registry_app
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[5]
MANIFESTS = (
    REPO_ROOT
    / "openbb_platform"
    / "extensions"
    / "agents"
    / "openbb_agents"
    / "notebook_guide"
    / "stories"
)
# TechTrade chapter 7's batch: three tickets, one failed sector gate.
BATCH = {
    "plan_id": "nb07-demo-2026",
    "verdict_gate_pass": True,
    "generated_at": "2026-01-05T15:00:00Z",
    "tickets": [
        {
            "symbol": "MSFT",
            "action": "Buy",
            "quantity": "50",
            "order_type": "Limit",
            "limit_price": "400.00",
            "account_masked": "***1234",
        },
        {
            "symbol": "AAPL",
            "action": "Buy",
            "quantity": "100",
            "order_type": "Limit",
            "limit_price": "180.00",
            "account_masked": "***1234",
        },
        {
            "symbol": "NVDA",
            "action": "Sell",
            "quantity": "25",
            "order_type": "Limit",
            "limit_price": "130.00",
            "account_masked": "***1234",
            "notes": "Take profit on rally",
        },
    ],
    "pricing": {"MSFT": "398.00", "AAPL": "179.50", "NVDA": "128.00"},
    "pre_execution_positions": {"MSFT": "0.05", "AAPL": "0.03", "NVDA": "0.12"},
    "verdict_gates": [
        {
            "name": "max_single_position",
            "threshold": "0.10",
            "actual": "0.08",
            "passed": True,
        },
        {
            "name": "max_sector_exposure",
            "threshold": "0.30",
            "actual": "0.35",
            "passed": False,
            "notes": "Tech cluster exceeds; reviewer sign-off required",
        },
    ],
    "generator_version": "nb07-1.0",
    "git_sha": "0000000000",
}


def _client(principal, store, root):
    app = compose_story_registry_app(
        FastAPI(), principal=principal, manifest_dir=MANIFESTS, store_path=store
    )
    app = compose_story_artifacts_app(
        app, principal=principal, store_path=store, artifact_root=root
    )
    app = compose_story_paper_batch_app(
        app, principal=principal, store_path=store, artifact_root=root
    )
    return TestClient(app)


@pytest.fixture
def harness(tmp_path):
    """Alice and Bob over one run store and one artifact root."""
    store = tmp_path / "runs.sqlite"
    root = tmp_path / "artifacts"
    alice = _client("alice", store, root)
    bob = _client("bob", store, root)

    def run(client=alice):
        response = client.post("/api/v1/story/runs", json={"story_id": "techtrade/07"})
        assert response.status_code == 200, response.text
        return response.json()["run_id"]

    def preview(batch=None, client=alice):
        return client.post("/api/v1/story/paper/batches/preview", json=batch or BATCH)

    def write(run_id, batch=None, client=alice):
        return client.post(
            f"/api/v1/story/runs/{run_id}/paper/batch_artifacts", json=batch or BATCH
        )

    def read(run_id, artifact_id, client=alice):
        response = client.get(f"/api/v1/story/runs/{run_id}/artifacts/{artifact_id}")
        assert response.status_code == 200, response.text
        return response.json()

    return SimpleNamespace(
        alice=alice,
        bob=bob,
        root=root,
        run=run,
        preview=preview,
        write=write,
        read=read,
    )


def _engine_batch():
    from openbb_techtrade.execution.order_sink import (
        OrderBatch,
        OrderTicket,
        PlanContext,
        VerdictGate,
    )

    return OrderBatch(
        tickets=tuple(
            OrderTicket(
                symbol=t["symbol"],
                action=t["action"],
                quantity=Decimal(t["quantity"]),
                order_type=t["order_type"],
                limit_price=Decimal(t["limit_price"]),
                account_masked=t["account_masked"],
                notes=t.get("notes", ""),
            )
            for t in BATCH["tickets"]
        ),
        plan_id=BATCH["plan_id"],
        verdict_gate_pass=True,
        generated_at=datetime(2026, 1, 5, 15, tzinfo=timezone.utc),
        plan_context=PlanContext(
            verdict_gates=tuple(VerdictGate(**g) for g in BATCH["verdict_gates"])
        ),
    )


def test_preview_returns_the_engines_content_sha(harness):
    """The batch SHA is the engine's own, over the executable tickets."""
    body = harness.preview().json()
    assert body["batch_sha256"] == _engine_batch().sha256()
    assert body["sha_short"] == body["batch_sha256"][:8]
    assert body["tickets"] == 3


def test_metadata_does_not_change_the_sha_but_tickets_do(harness):
    """Chapter 7: pricing and plan ID are not identity; a new ticket is."""
    base = harness.preview().json()["batch_sha256"]
    other_meta = {**BATCH, "plan_id": "different", "pricing": {"MSFT": "999"}}
    assert harness.preview(other_meta).json()["batch_sha256"] == base
    extra = {
        **BATCH,
        "tickets": [
            *BATCH["tickets"],
            {"symbol": "GOOGL", "action": "Buy", "quantity": "10"},
        ],
    }
    assert harness.preview(extra).json()["batch_sha256"] != base


def test_preview_surfaces_failed_gates_without_hiding_the_batch(harness):
    """A failed gate is reported for review; it is not silently dropped."""
    body = harness.preview().json()
    assert body["verdict_gate_pass"] is True
    assert body["failed_gates"] == ["max_sector_exposure"]
    assert body["executable"] is False


def test_preview_gives_the_operator_checklist(harness):
    """Chapter 9's 'type this at Fidelity' rows, one per ticket, in order."""
    rows = harness.preview().json()["checklist"]
    assert rows == [
        {
            "symbol": "MSFT",
            "action": "Buy",
            "quantity": "50",
            "order_type": "Limit",
            "limit_price": "400.00",
            "tif": "Day",
        },
        {
            "symbol": "AAPL",
            "action": "Buy",
            "quantity": "100",
            "order_type": "Limit",
            "limit_price": "180.00",
            "tif": "Day",
        },
        {
            "symbol": "NVDA",
            "action": "Sell",
            "quantity": "25",
            "order_type": "Limit",
            "limit_price": "130.00",
            "tif": "Day",
        },
    ]


@pytest.mark.parametrize(
    "ticket",
    [
        {"symbol": "MSFT", "action": "Buy", "quantity": "1", "notes": "=cmd|/C calc"},
        {"symbol": "msft!", "action": "Buy", "quantity": "1"},
        {"symbol": "MSFT", "action": "Buy", "quantity": "0"},
        {
            "symbol": "MSFT",
            "action": "Buy",
            "quantity": "1",
            "account_masked": "1234 5678",
        },
        {"symbol": "MSFT", "action": "Buy", "quantity": "1", "order_type": "Limit"},
    ],
)
def test_the_chapter_7_rejections_are_422s(harness, ticket):
    """Formula notes, bad symbols, zero quantity, raw accounts, missing limits."""
    response = harness.preview({**BATCH, "tickets": [ticket]})
    assert response.status_code == 422, response.text


def test_artifacts_are_the_engines_csv_and_workbook(harness):
    """The run gets the sink's exact Fidelity CSV and its six-sheet workbook."""
    run_id = harness.run()
    response = harness.write(run_id)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["batch_sha256"] == _engine_batch().sha256()
    csv_ref = body["csv"]["ref"]
    xlsx_ref = body["workbook"]["ref"]
    assert csv_ref["media_type"] == "text/csv"
    assert body["csv"]["name"] == f"fidelity-batch-{body['batch_sha256'][:8]}"
    text = harness.read(run_id, csv_ref["artifact_id"])["text"]
    assert text == (
        "MSFT,Buy,50,Limit,400,Day,***1234\r\n"
        "AAPL,Buy,100,Limit,180,Day,***1234\r\n"
        "NVDA,Sell,25,Limit,130,Day,***1234\r\n"
    )
    from openbb_techtrade.execution._xlsx_sheets import _SHEET_NAMES

    sheets = harness.read(run_id, xlsx_ref["artifact_id"])["workbook"]
    assert [s["name"] for s in sheets] == list(_SHEET_NAMES)


def test_the_workbook_shows_the_failed_gate_and_provenance(harness):
    """Plan Context lists every gate; Audit carries the full SHA and plan ID."""
    run_id = harness.run()
    body = harness.write(run_id).json()
    files = list(harness.root.rglob("*.xlsx"))
    assert len(files) == 1
    workbook = load_workbook(io.BytesIO(files[0].read_bytes()))
    context = list(workbook["Plan Context"].iter_rows(values_only=True))
    names = {row[0] for row in context if row and row[0]}
    assert {"max_single_position", "max_sector_exposure"} <= names
    audit = {
        row[0]: row[1] for row in workbook["Audit"].iter_rows(values_only=True) if row
    }
    assert body["batch_sha256"] in audit.values()
    assert "nb07-demo-2026" in audit.values()
    assert "nb07-1.0" in audit.values()
    # Pricing and pre-execution weights feed the Deviation and Concentration sheets.
    deviation = [
        c
        for row in workbook["Deviation Analysis"].iter_rows(values_only=True)
        for c in row
    ]
    assert 398 in deviation or 398.0 in deviation or "398.00" in deviation
    concentration = [
        c for row in workbook["Concentration"].iter_rows(values_only=True) for c in row
    ]
    assert any(c in (0.12, "0.12", 12.0) for c in concentration), concentration


def test_writing_the_same_batch_again_returns_the_same_artifacts(harness):
    """The sink is idempotent by batch SHA, and so is the run's copy."""
    run_id = harness.run()
    first = harness.write(run_id).json()
    second = harness.write(run_id).json()
    assert second["csv"]["ref"] == first["csv"]["ref"]
    assert second["workbook"]["ref"] == first["workbook"]["ref"]
    listed = harness.alice.get("/api/v1/story/artifacts", params={"run_id": run_id})
    assert len(listed.json()) == 2


def test_artifacts_need_a_run_the_principal_owns(harness):
    """Bob cannot write into Alice's run; unknown runs are 404."""
    run_id = harness.run()
    assert harness.write(run_id, client=harness.bob).status_code == 404
    assert harness.write("f" * 32).status_code == 404
    assert not list(harness.root.rglob("*.*"))


def test_batches_are_bounded(harness):
    """At most 50 tickets in a batch."""
    tickets = [{"symbol": "MSFT", "action": "Buy", "quantity": "1"}] * 51
    assert harness.preview({**BATCH, "tickets": tickets}).status_code == 422


def test_writing_artifacts_needs_no_approval_and_submits_nothing(harness):
    """Review documents are not paper actions: no request, no engine order."""
    run_id = harness.run()
    harness.write(run_id)
    with __import__("sqlite3").connect(
        harness.root.parent / "runs.sqlite"
    ) as connection:
        tables = {
            row[0]
            for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
    assert "story_paper_requests" not in tables
    assert not list(harness.root.parent.rglob("techtrade.sqlite"))


def test_responses_never_carry_a_filesystem_path(harness):
    """Neither the preview nor the artifacts name a local path."""
    run_id = harness.run()
    blob = json.dumps([harness.preview().json(), harness.write(run_id).json()])
    for marker in (
        str(harness.root),
        harness.root.as_posix(),
        "tmp",
        '.csv"',
        '.xlsx"',
    ):
        assert marker not in blob


def test_the_batch_routes_are_two_mcp_tools(tmp_path):
    """Preview and write; both authenticated, nothing approves or submits."""
    app = compose_story_paper_batch_app(
        FastAPI(), principal="alice", store_path=tmp_path / "r.sqlite"
    )
    routes = [r for r in app.routes if getattr(r, "openapi_extra", None)]
    assert {r.openapi_extra["mcp_config"]["name"] for r in routes} == {
        "story_paper_order_batch",
        "story_paper_batch_artifacts",
    }
    assert all(r.dependencies for r in routes)
    with pytest.raises(ValueError):
        compose_story_paper_batch_app(FastAPI(), principal="")


def test_unknown_fields_are_rejected(harness):
    """The batch schema is closed."""
    assert harness.preview({**BATCH, "broker": "live"}).status_code == 422
    assert paper_batch.MAX_TICKETS == 50
