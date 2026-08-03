"""T5 order sink — file-drop execution medium (#1719).

The techtrade "T5 execute" step turns a validated trading plan into
brokerage orders. Rather than route through a broker API (Alpaca / IBKR),
this implementation writes an **order batch** to disk as a paired
CSV + XLSX artifact. The operator uploads the CSV to Fidelity's Basket
Trading page; the XLSX is a human-review workbook that accompanies it.

Design shape (mirrors #1744's PortfolioStore Protocol):

- :class:`OrderSink` — ``runtime_checkable`` Protocol with three methods
  (``write_batch`` / ``list_batches`` / ``batch_exists``). Any concrete
  sink (paper simulator, Fidelity CSV, future Alpaca) satisfies it.
- :class:`PaperOrderSink` — the phase-1 backend. Writes CSV + XLSX under
  a caller-supplied output directory. Zero real-money surface — used
  by tests and dry runs.
- :class:`OrderTicket` / :class:`OrderBatch` / :class:`BatchArtifacts` —
  frozen dataclasses that model the batch and its written artifacts.

Idempotency: the batch's SHA256 is computed over the canonicalized order
tuple list (symbol/side/qty/limit/order_type/tif). Same plan → same SHA
→ ``batch_exists()`` returns the existing artifacts without rewriting.
The SHA is embedded in every filename so the CSV and XLSX are trivially
paired.

Atomicity: both files are written to a temp path in the same directory
then renamed. The operator never observes a half-written workbook.

Phase 1 (this module) ships only the ``Orders`` sheet in the XLSX. The
5 additional sheets (Batch Summary, Plan Context, Deviation Analysis,
Concentration, Audit) land in phase 2 as additive writes against the
frozen :class:`OrderSink` interface — no interface change needed.
"""

from __future__ import annotations

import contextlib
import csv
import hashlib
import logging
import os
import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Literal, Protocol, runtime_checkable

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Data types — frozen dataclasses so hash/equality are structural.
# ---------------------------------------------------------------------------


#: Fidelity Basket Trading valid actions.
Action = Literal["Buy", "Sell", "BuyToCover", "SellShort"]

#: Fidelity Basket Trading valid order types.
OrderType = Literal["Market", "Limit", "StopLoss", "StopLimit"]

#: Fidelity Basket Trading valid TIF codes.
Tif = Literal["Day", "GTC", "IOC", "FOK"]

_VALID_ACTIONS: frozenset[str] = frozenset({"Buy", "Sell", "BuyToCover", "SellShort"})
_VALID_ORDER_TYPES: frozenset[str] = frozenset(
    {"Market", "Limit", "StopLoss", "StopLimit"}
)
_VALID_TIFS: frozenset[str] = frozenset({"Day", "GTC", "IOC", "FOK"})

#: Symbol allowlist — same shape as the basket_id allowlist (#1714) plus
#: dot for share classes (BRK.B) and slash for warrants (WFC/WT).
_SYMBOL_RE = re.compile(r"^[A-Z][A-Z0-9./\-]{0,15}$")

#: CSV / Excel formula-injection leading chars. Any free-form string field
#: (notes, account_masked) starting with one of these is rejected loudly
#: at OrderTicket construction — Excel and Fidelity's basket importer
#: both interpret ``=SUM(...)`` / ``@cmd`` / ``+foo`` etc. as a formula
#: or command, so an operator (or an upstream planner) writing a
#: literal ``@`` into a note becomes remote-code-execution when the
#: basket lands on the reviewer's Excel. Loud rejection means a bad
#: note NEVER reaches either writer — one guard covers both formats.
_FORMULA_LEAD_CHARS = frozenset({"=", "+", "-", "@", "\t", "\r"})

#: Account-mask allowlist — expected shape is ``***1234``.
_ACCOUNT_MASK_RE = re.compile(r"^[*A-Za-z0-9_\-]{1,32}$")

#: Batch SHA short-form used in filenames (first N chars of the full SHA).
_SHA_SHORT_LEN = 8


@dataclass(frozen=True)
class OrderTicket:
    """One row in an order batch — the executable unit.

    Field order matches the Fidelity Basket Trading CSV column order so
    ``dataclasses.astuple`` on the canonical form yields the CSV row
    directly. See :func:`_canonicalize_ticket` for the SHA form.
    """

    symbol: str
    action: Action
    quantity: Decimal
    order_type: OrderType = "Market"
    limit_price: Decimal | None = None
    tif: Tif = "Day"
    account_masked: str | None = None  # last-4-mask, e.g. "***1234"
    notes: str = ""

    def __post_init__(self) -> None:
        """Validate ticket fields on construction; raise ValueError loudly."""
        # Validation runs on construction so a bad ticket never reaches
        # the writer. Loud errors, no silent coercion.
        if not _SYMBOL_RE.match(self.symbol):
            raise ValueError(
                f"OrderTicket.symbol must match {_SYMBOL_RE.pattern!r}; "
                f"got {self.symbol!r}"
            )
        if self.action not in _VALID_ACTIONS:
            raise ValueError(
                f"OrderTicket.action must be one of {sorted(_VALID_ACTIONS)}; "
                f"got {self.action!r}"
            )
        if self.order_type not in _VALID_ORDER_TYPES:
            raise ValueError(
                f"OrderTicket.order_type must be one of "
                f"{sorted(_VALID_ORDER_TYPES)}; got {self.order_type!r}"
            )
        if self.tif not in _VALID_TIFS:
            raise ValueError(
                f"OrderTicket.tif must be one of {sorted(_VALID_TIFS)}; "
                f"got {self.tif!r}"
            )
        if self.quantity <= 0:
            raise ValueError(
                f"OrderTicket.quantity must be positive; got {self.quantity}"
            )
        if self.order_type in ("Limit", "StopLimit") and self.limit_price is None:
            raise ValueError(
                f"OrderTicket.limit_price required for order_type "
                f"{self.order_type!r}"
            )
        if self.limit_price is not None and self.limit_price <= 0:
            raise ValueError(
                f"OrderTicket.limit_price must be positive; got " f"{self.limit_price}"
            )
        # CSV / Excel formula-injection guard on every free-form string
        # field. Symbol / action / order_type / tif are already allowlist-
        # validated above, so notes + account_masked are the only vectors.
        if self.notes and self.notes[0] in _FORMULA_LEAD_CHARS:
            raise ValueError(
                f"OrderTicket.notes must not start with a CSV/Excel "
                f"formula-injection character "
                f"{sorted(_FORMULA_LEAD_CHARS)}; got {self.notes!r}"
            )
        if self.account_masked is not None:
            if self.account_masked and self.account_masked[0] in _FORMULA_LEAD_CHARS:
                raise ValueError(
                    f"OrderTicket.account_masked must not start with a "
                    f"CSV/Excel formula-injection character "
                    f"{sorted(_FORMULA_LEAD_CHARS)}; got "
                    f"{self.account_masked!r}"
                )
            if not _ACCOUNT_MASK_RE.match(self.account_masked):
                raise ValueError(
                    f"OrderTicket.account_masked must match "
                    f"{_ACCOUNT_MASK_RE.pattern!r}; got "
                    f"{self.account_masked!r}"
                )


@dataclass(frozen=True)
class OrderBatch:
    """A validated set of order tickets destined for a broker.

    ``plan_id`` and ``verdict_gate_pass`` are provenance from the T4
    validation step; they surface in the XLSX audit sheet and the
    ``pi_order_batch`` audit row (phase 2).
    """

    tickets: tuple[OrderTicket, ...]
    plan_id: str = ""
    verdict_gate_pass: bool = False
    generated_at: datetime = field(default_factory=lambda: datetime.now(timezone.utc))

    def __post_init__(self) -> None:
        """Reject empty tickets and enforce tuple (stable hash)."""
        if not self.tickets:
            raise ValueError("OrderBatch.tickets must be non-empty")
        # Enforce tuple so hash is stable — dataclass frozen alone doesn't
        # forbid a list arg.
        if not isinstance(self.tickets, tuple):
            raise TypeError(
                f"OrderBatch.tickets must be a tuple; got "
                f"{type(self.tickets).__name__}"
            )

    def sha256(self) -> str:
        """Deterministic SHA256 over the canonical ticket sequence.

        Same tickets in the same order → same SHA. ``plan_id`` and
        ``generated_at`` are deliberately excluded — the SHA identifies
        the *executable content*, not the wall-clock of the write, so a
        re-run of the same plan yields idempotent artifacts.
        """
        h = hashlib.sha256()
        for t in self.tickets:
            h.update(_canonicalize_ticket(t).encode("utf-8"))
            h.update(b"\n")
        return h.hexdigest()

    def sha_short(self) -> str:
        """First 8 hex chars of :meth:`sha256` — used in filenames."""
        return self.sha256()[:_SHA_SHORT_LEN]


@dataclass(frozen=True)
class BatchArtifacts:
    """Paired paths produced by :meth:`OrderSink.write_batch`."""

    csv_path: Path
    xlsx_path: Path
    batch_sha256: str


@dataclass(frozen=True)
class BatchRecord:
    """Directory-listing summary of a previously-written batch.

    ``batch_sha256`` is the 8-char short form here (that's all the
    filename carries). The full 64-char SHA lives in the phase-2
    ``pi_order_batch`` audit row.
    """

    batch_sha256: str
    csv_path: Path
    xlsx_path: Path
    written_at: datetime


def _canonicalize_ticket(t: OrderTicket) -> str:
    """Pipe-delimited canonical form of a ticket for SHA hashing.

    Uses ``str(Decimal)`` so ``0.1`` and ``0.10`` are distinct — the
    Decimal round-trip is exact and matches how the user wrote it. Any
    ``None`` renders as the literal ``"None"`` so a limit-price
    presence/absence flips the SHA.
    """
    return "|".join(
        (
            t.symbol,
            t.action,
            str(t.quantity),
            t.order_type,
            str(t.limit_price) if t.limit_price is not None else "None",
            t.tif,
            t.account_masked or "",
        )
    )


# ---------------------------------------------------------------------------
# Protocol — the seam.
# ---------------------------------------------------------------------------


@runtime_checkable
class OrderSink(Protocol):
    """Structural contract every T5 execution backend satisfies.

    Mirrors the :class:`PortfolioStore` shape from #1744 — read/write/
    existence in three methods, no lifecycle noise.
    """

    def write_batch(self, batch: OrderBatch) -> BatchArtifacts:
        """Write both CSV and XLSX for ``batch``; return their paths.

        Idempotent: if a batch with the same SHA already exists in the
        sink, return its existing artifacts without rewriting (so a
        re-run of the same T5 step is safe).
        """
        ...  # pylint: disable=unnecessary-ellipsis

    def list_batches(self, since: date | None = None) -> list[BatchRecord]:
        """Enumerate previously-written batches, most recent first."""
        ...  # pylint: disable=unnecessary-ellipsis

    def batch_exists(self, batch_sha256: str) -> BatchArtifacts | None:
        """Return artifacts for ``batch_sha256`` if present, else None."""
        ...  # pylint: disable=unnecessary-ellipsis


class OrderSinkError(RuntimeError):
    """Raised for sink-level failures (path validation, atomic write)."""


# ---------------------------------------------------------------------------
# PaperOrderSink — phase-1 concrete backend.
# ---------------------------------------------------------------------------


class PaperOrderSink:
    """Writes CSV + XLSX to a caller-controlled output directory.

    No broker upload, no fills — this is the test/dry-run backend and
    the reference implementation of :class:`OrderSink`. The Fidelity
    CSV backend (phase 3) shares this class's CSV/XLSX writers and only
    overrides ``__init__`` to default the output directory and validate
    the path is outside the repo.
    """

    def __init__(self, output_dir: Path | str) -> None:
        self._output_dir = Path(output_dir).resolve()
        # Loud-empty: refuse to write to a nonexistent directory. Callers
        # are expected to have prepared it (mirrors the config layer
        # in the importer, #1744).
        if not self._output_dir.exists():
            raise OrderSinkError(
                f"output_dir does not exist: {self._output_dir!s} — "
                "create it before instantiating the sink"
            )
        if not self._output_dir.is_dir():
            raise OrderSinkError(
                f"output_dir must be a directory; got {self._output_dir!s}"
            )

    @property
    def output_dir(self) -> Path:
        """Directory where batches are written."""
        return self._output_dir

    # -- protocol methods ---------------------------------------------------

    def write_batch(self, batch: OrderBatch) -> BatchArtifacts:
        """Write CSV + XLSX for ``batch``; idempotent by batch SHA."""
        sha = batch.sha256()
        existing = self.batch_exists(sha)
        if existing is not None:
            logger.info(
                "PaperOrderSink: batch %s already written at %s — "
                "returning existing artifacts (idempotent)",
                sha[:_SHA_SHORT_LEN],
                existing.csv_path,
            )
            return existing

        stem = self._stem_for(batch)
        csv_path = self._output_dir / f"{stem}.csv"
        xlsx_path = self._output_dir / f"{stem}.xlsx"

        # Atomic writes: tmp + rename, both in the same directory so
        # rename is atomic on POSIX and near-atomic on NTFS.
        #
        # Write ORDER matters: XLSX (human-review workbook) is renamed
        # first, CSV (broker-upload file) is renamed last. A crash
        # between the two renames therefore never leaves an uploadable
        # CSV without its audit workbook — the operator's upload gate is
        # the xlsx-must-exist invariant that batch_exists() / list_batches()
        # enforce below.
        self._atomic_write(xlsx_path, lambda p: _write_xlsx(p, batch))
        self._atomic_write(csv_path, lambda p: _write_csv(p, batch))
        return BatchArtifacts(csv_path=csv_path, xlsx_path=xlsx_path, batch_sha256=sha)

    def list_batches(self, since: date | None = None) -> list[BatchRecord]:
        """List previously-written batches, most recent first.

        A batch is only listed when BOTH the CSV and its XLSX exist —
        write order (XLSX first, CSV last) makes a CSV-without-XLSX
        state impossible from a crashed write, so any partial batch
        found here is post-hoc tampering (hand-delete). Loud WARNING
        surfaces the anomaly instead of silently hiding it.
        """
        records: list[BatchRecord] = []
        for csv_path in sorted(self._output_dir.glob("*.csv"), reverse=True):
            sha = _sha_from_stem(csv_path.stem)
            if sha is None:
                continue
            xlsx_path = csv_path.with_suffix(".xlsx")
            if not xlsx_path.exists():
                logger.warning(
                    "PaperOrderSink: partial batch — csv %s exists but "
                    "xlsx %s is missing. Not a crash artifact (write order "
                    "would have prevented that); likely hand-delete. "
                    "Investigate before re-running the plan.",
                    csv_path.name,
                    xlsx_path.name,
                )
                continue
            written_at = datetime.fromtimestamp(
                csv_path.stat().st_mtime, tz=timezone.utc
            )
            if since is not None and written_at.date() < since:
                continue
            records.append(
                BatchRecord(
                    batch_sha256=sha,
                    csv_path=csv_path,
                    xlsx_path=xlsx_path,
                    written_at=written_at,
                )
            )
        return records

    def batch_exists(self, batch_sha256: str) -> BatchArtifacts | None:
        """Return existing artifacts for ``batch_sha256``, or None.

        Same partial-batch semantics as :meth:`list_batches`: an XLSX
        must be present for the batch to count as "written". A csv-only
        state triggers a WARNING because the write path can't produce
        one — write order renames XLSX first, CSV last.
        """
        short = batch_sha256[:_SHA_SHORT_LEN]
        # Filenames are `YYYY-MM-DD-<sha8>.{csv,xlsx}`. Scan by suffix.
        for csv_path in self._output_dir.glob(f"*-{short}.csv"):
            xlsx_path = csv_path.with_suffix(".xlsx")
            if xlsx_path.exists():
                return BatchArtifacts(
                    csv_path=csv_path,
                    xlsx_path=xlsx_path,
                    batch_sha256=batch_sha256,
                )
            logger.warning(
                "PaperOrderSink: batch_exists — csv %s present but xlsx "
                "%s missing. Treating as not-written; a re-run of the "
                "same plan will overwrite the csv. Post-hoc tampering?",
                csv_path.name,
                xlsx_path.name,
            )
        return None

    # -- helpers ------------------------------------------------------------

    def _stem_for(self, batch: OrderBatch) -> str:
        date_str = batch.generated_at.astimezone(timezone.utc).date().isoformat()
        return f"{date_str}-{batch.sha_short()}"

    @staticmethod
    def _atomic_write(final_path: Path, writer) -> None:
        tmp_path = final_path.with_name(final_path.name + ".tmp")
        try:
            writer(tmp_path)
        except Exception:
            # Clean up the partial temp so a retry starts clean.
            with contextlib.suppress(FileNotFoundError):
                tmp_path.unlink()
            raise
        os.replace(tmp_path, final_path)


def _sha_from_stem(stem: str) -> str | None:
    """Extract the sha_short from a batch filename stem.

    Filenames are ``YYYY-MM-DD-<sha8>``. The date prefix is fixed 10
    chars + hyphen + 8 hex chars. Return None if the stem doesn't match.
    """
    m = re.match(r"^\d{4}-\d{2}-\d{2}-([0-9a-f]{8})$", stem)
    return m.group(1) if m else None


# ---------------------------------------------------------------------------
# CSV writer — Fidelity Basket Trading format.
# ---------------------------------------------------------------------------


#: CSV column order per Fidelity Basket Trading spec. No header row is
#: written (Fidelity's importer expects raw data rows).
_CSV_COLUMNS = (
    "Symbol",
    "Action",
    "Quantity",
    "Order Type",
    "Limit Price",
    "TIF",
    "Account",
)


def _write_csv(path: Path, batch: OrderBatch) -> None:
    """Write ``batch`` to ``path`` in Fidelity Basket Trading CSV format.

    No header row — Fidelity's importer expects data rows only.
    Newline='' per csv module docs so no double-line-ending on Windows.
    """
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        for t in batch.tickets:
            writer.writerow(
                (
                    t.symbol,
                    t.action,
                    _fmt_decimal(t.quantity),
                    t.order_type,
                    _fmt_decimal(t.limit_price) if t.limit_price is not None else "",
                    t.tif,
                    t.account_masked or "",
                )
            )


def _fmt_decimal(d: Decimal) -> str:
    """Render a Decimal as a plain string with no trailing zero drift.

    ``normalize()`` collapses ``100.00`` to ``1E+2`` which some importers
    reject — so we render via ``str`` then strip trailing zeros only in
    the fractional part.
    """
    s = str(d)
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return s or "0"


# ---------------------------------------------------------------------------
# XLSX writer — phase-1 minimal (Orders sheet only).
# ---------------------------------------------------------------------------


def _write_xlsx(path: Path, batch: OrderBatch) -> None:
    """Write phase-1 minimal XLSX (Orders sheet only).

    Phase 2 will add: Batch Summary, Plan Context, Deviation Analysis,
    Concentration, Audit — plus embedded charts. Kept intentionally
    minimal here so the full write path (openpyxl workbook lifecycle,
    atomic rename, load-back verification) is exercised without the
    surface of 5 more sheets in this PR.
    """
    # Deferred import — openpyxl is a heavy dep and only the XLSX writer
    # needs it. If a downstream call site imports order_sink only for
    # the CSV path, no openpyxl load penalty.
    # pylint: disable=import-outside-toplevel
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"

    header = list(_CSV_COLUMNS) + ["Notes"]
    ws.append(header)
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    buy_fill = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
    sell_fill = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")

    for t in batch.tickets:
        row = [
            t.symbol,
            t.action,
            float(t.quantity),
            t.order_type,
            float(t.limit_price) if t.limit_price is not None else None,
            t.tif,
            t.account_masked or "",
            t.notes,
        ]
        ws.append(row)
        fill = buy_fill if t.action in ("Buy", "BuyToCover") else sell_fill
        for cell in ws[ws.max_row]:
            cell.fill = fill

    # Freeze the header row.
    ws.freeze_panes = "A2"

    # Column widths — approximate, no fancy auto-fit.
    widths = {"A": 12, "B": 12, "C": 12, "D": 12, "E": 14, "F": 8, "G": 14, "H": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    wb.save(path)


# ---------------------------------------------------------------------------
# Factory — env-var driven selection.
# ---------------------------------------------------------------------------


_ENV_SINK_KIND = "PI_ORDER_SINK"
_ENV_PAPER_DIR = "PI_ORDER_SINK_PAPER_DIR"


def get_default_sink(paper_dir: Path | str | None = None) -> OrderSink:
    """Return the configured T5 order sink.

    Selection rule:

    - ``PI_ORDER_SINK=paper`` (default) — returns a :class:`PaperOrderSink`
      rooted at ``paper_dir`` (arg), ``$PI_ORDER_SINK_PAPER_DIR`` (env), or
      a per-user default path (``~/.portfolio_intel/order_batches/``).
    - ``PI_ORDER_SINK=fidelity_csv`` — reserved for phase 3. Currently
      raises ``NotImplementedError`` with a clear message.
    - Anything else — loud ``ValueError``. No silent fallback.

    ``paper_dir`` is created if missing (unlike the raw :class:`PaperOrderSink`
    constructor, which refuses a nonexistent path — the factory offers a
    friendlier default for callers who don't want to think about it).
    """
    kind = os.environ.get(_ENV_SINK_KIND, "paper").strip().lower()
    if kind == "fidelity_csv":
        raise NotImplementedError(
            "PI_ORDER_SINK=fidelity_csv is reserved for phase 3 (#1719 P3). "
            "Use PI_ORDER_SINK=paper for now."
        )
    if kind != "paper":
        raise ValueError(
            f"PI_ORDER_SINK must be one of 'paper' | 'fidelity_csv'; " f"got {kind!r}"
        )

    resolved_dir: Path
    if paper_dir is not None:
        resolved_dir = Path(paper_dir)
    else:
        env_dir = os.environ.get(_ENV_PAPER_DIR)
        resolved_dir = (
            Path(env_dir)
            if env_dir
            else Path.home() / ".portfolio_intel" / "order_batches"
        )
    resolved_dir.mkdir(parents=True, exist_ok=True)
    logger.info(
        "get_default_sink: PI_ORDER_SINK=%s → PaperOrderSink at %s "
        "(operator uploads CSVs from here to Fidelity Basket Trading)",
        kind,
        resolved_dir,
    )
    return PaperOrderSink(resolved_dir)


# ---------------------------------------------------------------------------
# Convenience iterator so callers can build a batch from a plan of Orders
# without every phase having to reach into models.
# ---------------------------------------------------------------------------


def tickets_from_orders(
    orders: Iterable[object],
    account_masked: str | None = None,
) -> Iterator[OrderTicket]:
    """Adapt techtrade ``Order`` objects into :class:`OrderTicket` rows.

    Only entry legs are yielded — exits are simulated by the paper
    broker (``execution.broker``), not routed as separate orders here.
    Callers that need exits in the batch pass them through explicitly.

    Duck-typed on the ``Order`` fields the sink actually needs; keeps
    ``execution.broker.PaperBroker`` decoupled from the sink.

    Loud-drop discipline: an unknown ``side`` or missing required field
    logs a WARNING with the order's identity — silent drops here would
    turn a real signal (schema drift, new order type) into a truncated
    batch that looks complete. A batch that yielded zero tickets from a
    non-empty input logs an ERROR at the end.
    """
    side_to_action = {
        "buy": "Buy",
        "sell": "Sell",
        "buy_to_cover": "BuyToCover",
        "sell_short": "SellShort",
    }
    order_type_to_fidelity = {
        "market": "Market",
        "limit": "Limit",
        "stop": "StopLoss",
    }
    tif_to_fidelity = {"day": "Day", "gtc": "GTC"}

    input_count = 0
    yielded = 0
    for order in orders:
        input_count += 1
        symbol = getattr(order, "symbol", None)
        side = getattr(order, "side", None)
        intent = getattr(order, "intent", "entry")

        if side not in side_to_action:
            logger.warning(
                "tickets_from_orders: dropping order symbol=%r "
                "side=%r intent=%r — unknown side (schema drift?)",
                symbol,
                side,
                intent,
            )
            continue
        if intent != "entry":
            # exits handled by the paper broker, not the sink — silent
            # by design, but the intent=='exit_*' state is documented so
            # this is a design skip, not a signal drop.
            continue

        quantity = getattr(order, "quantity", None)
        if symbol is None or quantity is None:
            logger.warning(
                "tickets_from_orders: dropping order symbol=%r "
                "quantity=%r — required field missing",
                symbol,
                quantity,
            )
            continue

        try:
            ticket = OrderTicket(
                symbol=symbol,
                action=side_to_action[side],  # type: ignore[arg-type]
                quantity=quantity,
                order_type=order_type_to_fidelity.get(
                    getattr(order, "order_type", "market"), "Market"
                ),  # type: ignore[arg-type]
                limit_price=getattr(order, "limit_price", None),
                tif=tif_to_fidelity.get(
                    getattr(order, "tif", "day"), "Day"
                ),  # type: ignore[arg-type]
                account_masked=account_masked,
            )
        except ValueError as exc:
            logger.warning(
                "tickets_from_orders: dropping order symbol=%r — "
                "OrderTicket validation failed: %s",
                symbol,
                exc,
            )
            continue

        yielded += 1
        yield ticket

    if input_count > 0 and yielded == 0:
        logger.error(
            "tickets_from_orders: input had %d order(s) but yielded 0 "
            "tickets — all were dropped. Check WARNINGs above; the "
            "downstream OrderBatch construction will fail with "
            "'tickets must be non-empty'.",
            input_count,
        )
