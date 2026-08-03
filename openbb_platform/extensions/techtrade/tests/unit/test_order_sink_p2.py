"""P2 tests for :mod:`execution.order_sink` — 6-sheet XLSX (#1719).

Covers the P2 additive fields (``pricing``, ``pre_execution_positions``,
``plan_context``) and the 6-sheet workbook structure. P1 tests
(``test_order_sink.py``) still pass unchanged — this file only exercises
the new surface.

Every load-bearing assertion carries an R7.11 mutation-twin note.
"""

# ruff: noqa: D101, D102, D105

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openbb_techtrade.execution.order_sink import (
    OrderBatch,
    OrderTicket,
    PaperOrderSink,
    PlanContext,
    VerdictGate,
)


def _tk(
    symbol: str,
    qty: str = "10",
    action: str = "Buy",
    limit: str | None = None,
) -> OrderTicket:
    return OrderTicket(
        symbol=symbol,
        action=action,  # type: ignore[arg-type]
        quantity=Decimal(qty),
        order_type="Limit" if limit else "Market",
        limit_price=Decimal(limit) if limit else None,
    )


# ---------------------------------------------------------------------------
# Optional-field construction: OrderBatch accepts P2 fields, still SHA-stable
# ---------------------------------------------------------------------------


class TestP2AdditiveFields:
    def test_batch_accepts_new_optional_fields(self) -> None:
        b = OrderBatch(
            tickets=(_tk("MSFT", limit="400"),),
            pricing={"MSFT": Decimal("395")},
            pre_execution_positions={"AAPL": Decimal("0.3")},
            plan_context=PlanContext(
                verdict_gates=(
                    VerdictGate(
                        name="max_position",
                        threshold="0.10",
                        actual="0.08",
                        passed=True,
                    ),
                ),
                generator_version="1.0.0",
                git_sha="deadbeef",
            ),
        )
        assert b.pricing is not None
        assert b.pre_execution_positions is not None
        assert b.plan_context is not None

    def test_p2_fields_excluded_from_sha(self) -> None:
        """R7.11 twin: hashing pricing / plan_context would break P1
        callers because a re-priced re-run of the same plan would land
        at a new file. Verified — adding either to sha256() fails this
        test.
        """
        tickets = (_tk("MSFT", limit="400"),)
        b1 = OrderBatch(tickets=tickets)
        b2 = OrderBatch(
            tickets=tickets,
            pricing={"MSFT": Decimal("395")},
            pre_execution_positions={"MSFT": Decimal("0.10")},
            plan_context=PlanContext(
                verdict_gates=(),
                generator_version="test",
                git_sha="cafebabe",
            ),
        )
        assert b1.sha256() == b2.sha256(), (
            "P2 optional fields must NOT feed the batch SHA — "
            "idempotency invariant depends on it"
        )


# ---------------------------------------------------------------------------
# 6-sheet structure — every named sheet present in tab order
# ---------------------------------------------------------------------------


_EXPECTED_SHEETS = (
    "Orders",
    "Batch Summary",
    "Plan Context",
    "Deviation Analysis",
    "Concentration",
    "Audit",
)


class TestSixSheetStructure:
    def test_six_sheets_present_in_tab_order(self, tmp_path: Path) -> None:
        """R7.11 twin: dropping any _write_<sheet>_sheet call in
        _write_xlsx breaks this exact-sequence assertion.
        """
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(tickets=(_tk("MSFT", limit="400"),))
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        assert (
            tuple(wb.sheetnames) == _EXPECTED_SHEETS
        ), f"expected {_EXPECTED_SHEETS!r}; got {tuple(wb.sheetnames)!r}"

    def test_p1_style_batch_still_writes_valid_workbook(self, tmp_path: Path) -> None:
        """A batch with no P2 optional fields (P1-style) still produces
        a loadable 6-sheet workbook. Placeholders on sheets 3-5.

        R7.11 twin: any of the sheet writers raising on None optional
        fields would fail write_batch here.
        """
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(tickets=(_tk("MSFT", limit="400"),))
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        # Placeholder markers on the sheets that need optional fields.
        assert "not provided" in str(wb["Plan Context"].cell(row=2, column=1).value)
        assert "not provided" in str(
            wb["Deviation Analysis"].cell(row=2, column=1).value
        )
        assert "not provided" in str(wb["Concentration"].cell(row=2, column=1).value)


# ---------------------------------------------------------------------------
# Batch Summary — verdict-gate loud-red on failure
# ---------------------------------------------------------------------------


class TestBatchSummarySheet:
    def test_verdict_gate_pass_renders_green(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="400"),),
            verdict_gate_pass=True,
        )
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Batch Summary"]
        gate_cell = ws.cell(row=5, column=2)
        assert gate_cell.value == "PASS"

    def test_verdict_gate_fail_gets_loud_red_fill(self, tmp_path: Path) -> None:
        """R7.11 twin: reverting the `if not batch.verdict_gate_pass`
        red-fill branch removes the fill and this fill-color assertion
        fails.
        """
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="400"),),
            verdict_gate_pass=False,
        )
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Batch Summary"]
        gate_row_cell = ws.cell(row=5, column=1)
        # openpyxl stores hex as "00" + our 6-char code
        assert "FADBD8" in (gate_row_cell.fill.start_color.rgb or "").upper()

    def test_sha_appears_in_summary(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(tickets=(_tk("MSFT", limit="400"),))
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Batch Summary"]
        assert ws.cell(row=2, column=2).value == batch.sha_short()
        assert ws.cell(row=3, column=2).value == batch.sha256()


# ---------------------------------------------------------------------------
# Plan Context — one row per gate, failed gates red
# ---------------------------------------------------------------------------


class TestPlanContextSheet:
    def test_gates_render_one_per_row(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="400"),),
            plan_context=PlanContext(
                verdict_gates=(
                    VerdictGate(
                        name="max_pos",
                        threshold="0.10",
                        actual="0.08",
                        passed=True,
                    ),
                    VerdictGate(
                        name="max_sector",
                        threshold="0.30",
                        actual="0.35",
                        passed=False,
                    ),
                ),
            ),
        )
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Plan Context"]
        assert ws.cell(row=2, column=1).value == "max_pos"
        assert ws.cell(row=3, column=1).value == "max_sector"
        assert ws.cell(row=2, column=4).value == "PASS"
        assert ws.cell(row=3, column=4).value == "FAIL"

    def test_failed_gate_gets_red_fill(self, tmp_path: Path) -> None:
        """R7.11 twin: dropping the `if not g.passed` branch removes the
        fill and this assertion fails.
        """
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="400"),),
            plan_context=PlanContext(
                verdict_gates=(
                    VerdictGate(
                        name="max_sector",
                        threshold="0.30",
                        actual="0.35",
                        passed=False,
                    ),
                ),
            ),
        )
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Plan Context"]
        # row 2 (only gate); check first cell of the row for red fill
        assert "FADBD8" in (ws.cell(row=2, column=1).fill.start_color.rgb or "").upper()


# ---------------------------------------------------------------------------
# Deviation Analysis — table + embedded bar chart
# ---------------------------------------------------------------------------


class TestDeviationSheet:
    def test_deviation_computed_from_pricing(self, tmp_path: Path) -> None:
        """MSFT limit=410, last_close=400 → deviation = +2.500%.

        R7.11 twin: if the formula reversed to `(last_close - limit) /
        last_close`, the result would be -2.5%. Verified by flipping —
        this test fails.
        """
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="410"),),
            pricing={"MSFT": Decimal("400")},
        )
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Deviation Analysis"]
        dev = ws.cell(row=2, column=6).value
        assert dev == pytest.approx(2.500, abs=1e-3)

    def test_large_deviation_gets_warn_fill(self, tmp_path: Path) -> None:
        """R7.11 twin: raising the 5% threshold to 10% would let 6%
        deviations through unhighlighted. Verified.
        """
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        # 6% deviation: limit=424, last=400
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="424"),),
            pricing={"MSFT": Decimal("400")},
        )
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Deviation Analysis"]
        # amber warn fill on row 2
        assert "FFF4E5" in (ws.cell(row=2, column=1).fill.start_color.rgb or "").upper()

    def test_bar_chart_embedded_when_data_present(self, tmp_path: Path) -> None:
        """R7.11 twin: skipping the ws.add_chart() call removes the chart
        and _charts becomes empty. Verified.
        """
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="400"), _tk("AAPL", limit="180")),
            pricing={"MSFT": Decimal("395"), "AAPL": Decimal("175")},
        )
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Deviation Analysis"]
        assert len(ws._charts) == 1  # noqa: SLF001

    def test_no_chart_when_pricing_absent(self, tmp_path: Path) -> None:
        """Placeholder-row path: no data → no chart, no crash."""
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(tickets=(_tk("MSFT", limit="400"),))
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Deviation Analysis"]
        assert len(ws._charts) == 0  # noqa: SLF001


# ---------------------------------------------------------------------------
# Concentration — table math + two pie charts + loud red on > 10%
# ---------------------------------------------------------------------------


class TestConcentrationSheet:
    def test_pre_weight_appears_when_provided(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="400"),),
            pre_execution_positions={"MSFT": Decimal("0.05"), "AAPL": Decimal("0.30")},
        )
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Concentration"]
        # Alphabetized rows: AAPL row 2, MSFT row 3
        assert ws.cell(row=2, column=1).value == "AAPL"
        assert ws.cell(row=3, column=1).value == "MSFT"

    def test_over_concentration_row_gets_red_fill(self, tmp_path: Path) -> None:
        """AAPL is 30% pre-exec — over the 10% single-position limit.

        R7.11 twin: raising the threshold from 0.10 to 0.50 lets this
        pass through unhighlighted. Verified.
        """
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="400"),),
            pre_execution_positions={"AAPL": Decimal("0.30")},
        )
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Concentration"]
        # AAPL row → row 2, red fill
        assert "FADBD8" in (ws.cell(row=2, column=1).fill.start_color.rgb or "").upper()

    def test_two_pie_charts_embedded(self, tmp_path: Path) -> None:
        """R7.11 twin: dropping either add_chart() call halves the count."""
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="400"),),
            pre_execution_positions={"AAPL": Decimal("0.3")},
        )
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Concentration"]
        assert len(ws._charts) == 2  # noqa: SLF001


# ---------------------------------------------------------------------------
# Audit sheet — provenance fields present
# ---------------------------------------------------------------------------


class TestAuditSheet:
    def test_audit_carries_full_sha_and_context(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="400"),),
            plan_id="plan-xyz",
            verdict_gate_pass=True,
            plan_context=PlanContext(
                generator_version="techtrade-1.2.3",
                git_sha="abc123def456",
            ),
        )
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Audit"]
        # Build a dict from the sheet rows for easier assertion.
        rows = {
            ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
            for r in range(2, ws.max_row + 1)
        }
        assert rows["Batch SHA (full)"] == batch.sha256()
        assert rows["Plan ID"] == "plan-xyz"
        assert rows["Generator version"] == "techtrade-1.2.3"
        assert rows["Git SHA"] == "abc123def456"

    def test_audit_placeholders_when_context_absent(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(tickets=(_tk("MSFT", limit="400"),))
        art = sink.write_batch(batch)
        wb = load_workbook(art.xlsx_path)
        ws = wb["Audit"]
        rows = {
            ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
            for r in range(2, ws.max_row + 1)
        }
        assert rows["Generator version"] == "not provided"
        assert rows["Git SHA"] == "not provided"


# ---------------------------------------------------------------------------
# Idempotency preserved — same P2-rich batch, second write is a no-op
# ---------------------------------------------------------------------------


class TestP2IdempotencyPreserved:
    def test_second_write_of_p2_batch_is_no_op(self, tmp_path: Path) -> None:
        """R7.11 twin: if _write_xlsx or _write_csv touched the file
        again despite the sha match, mtime would differ. Verified — this
        is the same load-bearing invariant P1 established, re-verified
        against the fuller P2 batch.
        """
        sink = PaperOrderSink(tmp_path)
        batch = OrderBatch(
            tickets=(_tk("MSFT", limit="400"),),
            pricing={"MSFT": Decimal("395")},
            pre_execution_positions={"AAPL": Decimal("0.3")},
            plan_context=PlanContext(
                verdict_gates=(
                    VerdictGate(
                        name="max_pos",
                        threshold="0.10",
                        actual="0.08",
                        passed=True,
                    ),
                ),
            ),
        )
        a1 = sink.write_batch(batch)
        mtime1 = a1.xlsx_path.stat().st_mtime_ns
        a2 = sink.write_batch(batch)
        assert a1.xlsx_path == a2.xlsx_path
        assert mtime1 == a2.xlsx_path.stat().st_mtime_ns
